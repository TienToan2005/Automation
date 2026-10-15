"""Báo cáo ra vào: số liệu tổng hợp, xuất CSV / Excel.

Mọi "ngày" và "giờ" trong báo cáo là giờ địa phương (LAB_UTC_OFFSET_HOURS, mặc định UTC+7); DB vẫn lưu UTC.
Cách đếm: một lượt quét thẻ + mặt chỉ tính MỘT lần, ở dòng `face_verify` (dòng `card_scan` "pending" là lúc bắt đầu phiên).
Thẻ bị từ chối ngay ở bước thẻ (thẻ lạ, thu hồi, khóa, chưa đăng ký mặt, AI hỏng) không có dòng `face_verify`
nên được tính từ dòng `card_scan` bị từ chối.
"""
import csv
import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

from fastapi import HTTPException
from sqlalchemy.orm import Session

from ..core.clock import local_day_start_utc, local_today, to_local, utcnow
from ..core.config import settings
from ..core.database import SessionLocal
from ..models import AccessLog, AccessResult, Alert, LogEvent, User

MAX_DAYS = 366
DEFAULT_DAYS = 7
XLSX_MAX_LOG_ROWS = 20_000          # openpyxl giữ cả sổ trong RAM; nhật ký lớn hơn thì dùng CSV
TOP_N = 10

EVENT_LABEL = {"card_scan": "Quẹt thẻ", "face_verify": "Quét mặt", "exit_button": "Nút trong", "remote_open": "Mở từ xa",
               "admin_action": "Quản trị"}
RESULT_LABEL = {"granted": "Cho vào", "denied": "Từ chối", "pending": "Đang xác thực", "timeout": "Hết giờ",
                "cancelled": "Đã hủy"}
REASON_LABEL = {
    "unknown_card": "Thẻ không tồn tại", "card_revoked": "Thẻ đã thu hồi", "account_suspended": "Tài khoản bị khóa",
    "not_enrolled": "Chưa đăng ký khuôn mặt", "ai_unavailable": "AI không khả dụng", "face_mismatch": "Sai khuôn mặt",
    "spoof_pad": "Giả mạo (ảnh in / màn hình)", "challenge_failed": "Không qua thử thách người thật",
    "timeout": "Hết giờ quét mặt", "pad_unavailable": "Thiếu model chống giả mạo",
}
ALERT_REASON_LABEL = {**REASON_LABEL, "revoked_card_used": "Dùng thẻ đã thu hồi"}
LEVEL_LABEL = {"high": "Cao", "medium": "Trung bình"}
_DURATION = re.compile(r"^(\d+) ms")        # AccessLog.note của dòng face_verify bắt đầu bằng "<số> ms"


# ------------------------------ Khoảng thời gian ------------------------------
def parse_range(date_from: str | None, date_to: str | None) -> tuple[date, date]:
    """Hai đầu mút là NGÀY địa phương (YYYY-MM-DD), gồm cả hai đầu. Mặc định: 7 ngày gần nhất tính đến hôm nay."""
    def parse(value: str, name: str) -> date:
        try:
            return date.fromisoformat(value)
        except ValueError:
            raise HTTPException(400, f"{name} không hợp lệ, cần dạng YYYY-MM-DD: {value}")

    end = parse(date_to, "date_to") if date_to else local_today()
    start = parse(date_from, "date_from") if date_from else end - timedelta(days=DEFAULT_DAYS - 1)
    if start > end:
        raise HTTPException(400, "Ngày bắt đầu phải trước hoặc bằng ngày kết thúc")
    if (end - start).days + 1 > MAX_DAYS:
        raise HTTPException(400, f"Khoảng thời gian tối đa {MAX_DAYS} ngày")
    return start, end


def utc_window(start: date, end: date) -> tuple[datetime, datetime]:
    """[từ, đến) theo UTC tương ứng với các ngày địa phương start..end."""
    return local_day_start_utc(start), local_day_start_utc(end + timedelta(days=1))


def outcome(event: str, result: str) -> str | None:
    """Một dòng nhật ký tính vào loại nào: entry | denied | timeout | cancelled | exit_button | remote_open | admin_action."""
    if event == LogEvent.FACE_VERIFY:
        return {AccessResult.GRANTED: "entry", AccessResult.DENIED: "denied", AccessResult.TIMEOUT: "timeout",
                AccessResult.CANCELLED: "cancelled"}.get(result)
    if event == LogEvent.CARD_SCAN:
        return "denied" if result == AccessResult.DENIED else None
    return {LogEvent.EXIT_BUTTON: "exit_button", LogEvent.REMOTE_OPEN: "remote_open",
            LogEvent.ADMIN_ACTION: "admin_action"}.get(event)


def _mean(values: list[float], digits: int) -> float | None:
    return round(sum(values) / len(values), digits) if values else None


# ------------------------------ Tổng hợp ------------------------------
def build_summary(db: Session, start: date, end: date, top: int | None = TOP_N) -> dict:
    """`top` = số sinh viên tối đa trong các bảng xếp hạng (None = tất cả)."""
    lo, hi = utc_window(start, end)
    days = {start + timedelta(days=i): Counter() for i in range((end - start).days + 1)}
    hours = [Counter() for _ in range(24)]
    totals, reasons, per_user = Counter(), Counter(), defaultdict(Counter)
    confidence, pad, duration = [], [], []

    rows = (db.query(AccessLog.ts, AccessLog.event, AccessLog.student_id, AccessLog.result, AccessLog.reason,
                     AccessLog.confidence, AccessLog.pad_score, AccessLog.note)
            .filter(AccessLog.ts >= lo, AccessLog.ts < hi).yield_per(5000))
    for ts, event, student_id, result, reason, conf, pad_score, note in rows:
        kind = outcome(event, result)
        if kind is None:
            continue
        local = to_local(ts)
        totals[kind] += 1
        if local.date() in days:
            days[local.date()][kind] += 1
        if kind == "entry":
            hours[local.hour]["entry"] += 1
            per_user[student_id]["entry"] += 1
            if conf is not None:
                confidence.append(conf)
            if pad_score is not None:
                pad.append(pad_score)
            if m := _DURATION.match(note or ""):
                duration.append(int(m.group(1)))
        elif kind in ("denied", "timeout"):
            hours[local.hour]["denied"] += 1
            reasons[reason or kind] += 1
            per_user[student_id]["denied"] += 1
    per_user.pop(None, None)

    decided = totals["entry"] + totals["denied"] + totals["timeout"]
    names = dict(db.query(User.student_id, User.full_name).filter(User.student_id.in_(list(per_user))).all()) if per_user else {}
    users = [{"student_id": sid, "full_name": names.get(sid, ""), "entries": c["entry"], "denied": c["denied"]}
             for sid, c in per_user.items()]

    alerts = db.query(Alert).filter(Alert.ts >= lo, Alert.ts < hi).all()
    resolved = [a for a in alerts if a.resolved]
    minutes = [(a.resolved_at - a.ts).total_seconds() / 60 for a in resolved if a.resolved_at]

    return {
        "range": {"from": start.isoformat(), "to": end.isoformat(), "days": len(days), "today": local_today().isoformat(),
                  "utc_offset_hours": settings.utc_offset_hours, "generated_at": to_local(utcnow()).strftime("%Y-%m-%d %H:%M:%S")},
        "totals": {
            "entries": totals["entry"], "denied": totals["denied"], "timeouts": totals["timeout"],
            "cancelled": totals["cancelled"], "exit_button": totals["exit_button"], "remote_open": totals["remote_open"],
            "admin_actions": totals["admin_action"],
            "success_rate": round(totals["entry"] / decided, 4) if decided else None,
            "unique_users": sum(1 for c in per_user.values() if c["entry"]),
            "avg_confidence": _mean(confidence, 3), "avg_pad_score": _mean(pad, 3), "avg_duration_ms": _mean(duration, 0),
        },
        "by_day": [{"date": d.isoformat(), "entries": c["entry"], "denied": c["denied"], "timeouts": c["timeout"],
                    "exit_button": c["exit_button"], "remote_open": c["remote_open"]} for d, c in days.items()],
        "by_hour": [{"hour": h, "entries": c["entry"], "denied": c["denied"]} for h, c in enumerate(hours)],
        "denial_reasons": [{"reason": r, "label": REASON_LABEL.get(r, r), "count": n} for r, n in reasons.most_common()],
        "top_users": sorted((u for u in users if u["entries"]), key=lambda u: (-u["entries"], u["student_id"]))[:top],
        "top_denied": sorted((u for u in users if u["denied"]), key=lambda u: (-u["denied"], u["student_id"]))[:top],
        "alerts": {
            "total": len(alerts), "open": len(alerts) - len(resolved), "resolved": len(resolved),
            "avg_resolve_minutes": _mean(minutes, 1),
            "by_reason": [{"reason": r, "label": ALERT_REASON_LABEL.get(r, r), "count": n}
                          for r, n in Counter(a.reason for a in alerts).most_common()],
            "by_level": {lv: sum(1 for a in alerts if a.level == lv) for lv in ("high", "medium")},
        },
    }


# ------------------------------ CSV ------------------------------
def _safe(value):
    """Chống "CSV injection": ô bắt đầu bằng = + - @ bị Excel hiểu là công thức. Dữ liệu như `note` do người dùng nhập."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _csv_line(values) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerow([_safe(v) for v in values])
    return buf.getvalue()


def _local_text(ts: datetime | None) -> str:
    return to_local(ts).strftime("%Y-%m-%d %H:%M:%S") if ts else ""


def log_query(db: Session, lo: datetime, hi: datetime, *, result=None, event=None, reason=None, student_id=None):
    q = db.query(AccessLog).filter(AccessLog.ts >= lo, AccessLog.ts < hi)
    for column, value in ((AccessLog.result, result), (AccessLog.event, event), (AccessLog.reason, reason),
                          (AccessLog.student_id, (student_id or "").strip() or None)):
        if value:
            q = q.filter(column == value)
    return q


def _log_row(a: AccessLog, names: dict) -> list:
    return [_local_text(a.ts), a.student_id or "", names.get(a.student_id, ""), EVENT_LABEL.get(a.event, a.event),
            RESULT_LABEL.get(a.result, a.result), REASON_LABEL.get(a.reason, a.reason), a.reason,
            "" if a.confidence is None else round(a.confidence, 3), "" if a.pad_score is None else round(a.pad_score, 3),
            a.device, a.session_id or "", a.note]


def log_columns() -> list[str]:
    return [f"Thời gian (UTC{settings.utc_offset_hours:+d})", "Mã SV", "Họ tên", "Sự kiện", "Kết quả", "Lý do", "Mã lý do",
            "Độ khớp", "Điểm chống giả", "Thiết bị", "Phiên", "Ghi chú"]


def iter_logs_csv(start: date, end: date, **filters):
    """Sinh từng dòng CSV (UTF-8 có BOM để Excel đọc đúng tiếng Việt). Tự mở phiên DB vì chạy sau khi request trả về."""
    lo, hi = utc_window(start, end)
    yield "﻿" + _csv_line(log_columns())
    with SessionLocal() as db:
        names = dict(db.query(User.student_id, User.full_name).all())
        q = log_query(db, lo, hi, **filters).order_by(AccessLog.ts).limit(settings.export_max_rows)
        for a in q.yield_per(2000):
            yield _csv_line(_log_row(a, names))


def alert_columns() -> list[str]:
    return [f"Thời gian (UTC{settings.utc_offset_hours:+d})", "Mức", "Loại", "Mã SV", "Họ tên", "Nội dung", "Trạng thái",
            "Người xử lý", "Thời gian xử lý", "Ghi chú xử lý", "Có ảnh"]


def _alert_row(a: Alert, names: dict) -> list:
    return [_local_text(a.ts), LEVEL_LABEL.get(a.level, a.level), ALERT_REASON_LABEL.get(a.reason, a.reason), a.student_id or "",
            names.get(a.student_id, ""), a.message, "Đã xử lý" if a.resolved else "Chưa xử lý", a.resolved_by or "",
            _local_text(a.resolved_at), a.resolution_note or "", "Có" if a.snapshot else ""]


def iter_alerts_csv(start: date, end: date):
    lo, hi = utc_window(start, end)
    yield "﻿" + _csv_line(alert_columns())
    with SessionLocal() as db:
        names = dict(db.query(User.student_id, User.full_name).all())
        for a in db.query(Alert).filter(Alert.ts >= lo, Alert.ts < hi).order_by(Alert.ts).yield_per(2000):
            yield _csv_line(_alert_row(a, names))


# ------------------------------ Excel ------------------------------
def build_workbook(db: Session, start: date, end: date) -> bytes:
    """Sổ Excel nhiều trang (có biểu đồ). Ném ImportError nếu chưa cài openpyxl."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    s = build_summary(db, start, end, top=None)
    lo, hi = utc_window(start, end)
    names = dict(db.query(User.student_id, User.full_name).all())
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")

    def sheet(title: str, columns: list[str], rows: list[list], widths: list[int] | None = None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        for cells in ws.iter_rows(min_row=2):      # chuỗi bắt đầu bằng "=" (tên, ghi chú do người dùng nhập) không được thành công thức
            for cell in cells:
                if cell.data_type == "f":
                    cell.data_type = "s"
        for i in range(len(columns)):
            ws.column_dimensions[get_column_letter(i + 1)].width = (widths or [])[i] if widths and i < len(widths) else 16
        ws.freeze_panes = "A2"
        return ws

    t, r = s["totals"], s["range"]
    pct = lambda v: "" if v is None else f"{v * 100:.1f}%"
    sheet("Tổng quan", ["Chỉ số", "Giá trị"], [
        ["Khoảng thời gian", f"{r['from']} → {r['to']} ({r['days']} ngày, UTC{r['utc_offset_hours']:+d})"],
        ["Lập lúc", r["generated_at"]],
        ["Lượt vào hợp lệ (thẻ + khuôn mặt)", t["entries"]], ["Lượt bị từ chối", t["denied"]], ["Hết giờ quét mặt", t["timeouts"]],
        ["Tỉ lệ thành công", pct(t["success_rate"])], ["Số sinh viên khác nhau đã vào", t["unique_users"]],
        ["Mở bằng nút trong", t["exit_button"]], ["Mở từ xa (admin)", t["remote_open"]],
        ["Độ khớp khuôn mặt trung bình", "" if t["avg_confidence"] is None else t["avg_confidence"]],
        ["Điểm chống giả mạo trung bình", "" if t["avg_pad_score"] is None else t["avg_pad_score"]],
        ["Thời gian xác thực trung bình (ms)", "" if t["avg_duration_ms"] is None else t["avg_duration_ms"]],
        ["Cảnh báo an ninh", s["alerts"]["total"]], ["  chưa xử lý", s["alerts"]["open"]], ["  đã xử lý", s["alerts"]["resolved"]],
        ["  thời gian xử lý trung bình (phút)", "" if s["alerts"]["avg_resolve_minutes"] is None else s["alerts"]["avg_resolve_minutes"]],
    ], [42, 44], first=True)

    n = len(s["by_day"])
    ws = sheet("Theo ngày", ["Ngày", "Cho vào", "Từ chối", "Hết giờ", "Nút trong", "Mở từ xa"],
               [[d["date"], d["entries"], d["denied"], d["timeouts"], d["exit_button"], d["remote_open"]] for d in s["by_day"]])
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title, chart.height, chart.width = "Lượt ra vào theo ngày", 8, 22
    chart.add_data(Reference(ws, min_col=2, max_col=6, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    ws.add_chart(chart, "H2")

    ws = sheet("Theo giờ", ["Giờ", "Cho vào", "Từ chối + hết giờ"], [[f"{h['hour']:02d}h", h["entries"], h["denied"]] for h in s["by_hour"]])
    chart = BarChart()
    chart.type, chart.title, chart.height, chart.width = "col", "Phân bố theo giờ trong ngày", 8, 22
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=25), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=25))
    ws.add_chart(chart, "E2")

    sheet("Lý do từ chối", ["Lý do", "Mã", "Số lượt"], [[d["label"], d["reason"], d["count"]] for d in s["denial_reasons"]], [36, 20, 12])
    sheet("Sinh viên", ["Mã SV", "Họ tên", "Lượt vào", "Lượt từ chối"],
          [[u["student_id"], u["full_name"], u["entries"], u["denied"]]
           for u in {x["student_id"]: x for x in s["top_users"] + s["top_denied"]}.values()], [16, 30, 12, 14])

    alerts = db.query(Alert).filter(Alert.ts >= lo, Alert.ts < hi).order_by(Alert.ts).all()
    sheet("Cảnh báo", alert_columns(), [_alert_row(a, names) for a in alerts], [20, 12, 30, 14, 24, 60, 14, 14, 20, 40, 8])

    logs = log_query(db, lo, hi).order_by(AccessLog.ts).limit(XLSX_MAX_LOG_ROWS + 1).all()
    ws = sheet("Nhật ký", log_columns(), [_log_row(a, names) for a in logs[:XLSX_MAX_LOG_ROWS]],
               [20, 14, 24, 12, 14, 30, 18, 10, 14, 18, 14, 40])
    if len(logs) > XLSX_MAX_LOG_ROWS:
        ws.append([f"… đã cắt ở {XLSX_MAX_LOG_ROWS} dòng đầu; dùng xuất CSV để lấy đủ."])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
